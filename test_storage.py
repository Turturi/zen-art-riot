"""Run with: python test_storage.py"""
import os
import tempfile
from openpyxl import load_workbook

with tempfile.TemporaryDirectory() as directory:
    os.environ['EXCEL_PATH'] = os.path.join(directory, 'registrations.xlsx')
    from storage import ExcelStore, EVENTS
    store = ExcelStore()
    base = dict(name='Test Student', mobile='9876543210', college='Test College', department='Computer Science', year='1st Year')
    ids = []
    for event in ['Singing', 'Singing', 'Reel Roulette', 'Dance']:
        reg_id, error = store.register({**base, 'event': event})
        assert not error, error
        ids.append(reg_id)
    assert len(set(ids)) == 4
    wb = load_workbook(os.environ['EXCEL_PATH'])
    assert set(EVENTS).issubset(set(wb.sheetnames))
    assert wb['Singing'].max_row == 3
    assert wb['Reel Roulette'].max_row == 2
    assert wb['Dance'].max_row == 2
    assert len(store.rows()) == 4
    print('PASS: same mobile can register repeatedly, event tabs are separate, unique IDs created.')
