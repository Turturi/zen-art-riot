"""Excel workbook with one worksheet per event; local or Azure Blob persistence."""
import io
import os
import re
import threading
import time
import uuid
from datetime import datetime
from zoneinfo import ZoneInfo
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

EVENTS = ["Singing", "Reel Roulette", "Film Making", "Dance", "Antakshari", "Mr Sau Ms Sau", "Star Night", "Fashion Show"]
HEADERS = ["Registration ID", "Student Name", "Mobile Number", "Department", "Year", "College", "Event", "Registration Date", "Registration Time", "Category", "Batch"]
LOCAL_PATH = os.path.abspath(os.getenv("EXCEL_PATH", os.path.join(os.path.dirname(__file__), "college_fest_registrations.xlsx")))
_LOCK = threading.RLock()

def _new_workbook():
    wb = Workbook()
    wb.remove(wb.active)
    for event in EVENTS:
        _ensure_sheet(wb, event)
    return wb

def _ensure_sheet(wb, event):
    if event not in wb.sheetnames:
        ws = wb.create_sheet(event)
        ws.append(HEADERS)
    else:
        ws = wb[event]
        # Preserve old registrations and extend existing sheets with new fields.
        if ws.max_row == 0 or ws.cell(1, 1).value is None:
            ws.append(HEADERS)
        existing = [ws.cell(1, i).value for i in range(1, ws.max_column + 1)]
        for header in HEADERS:
            if header not in existing:
                ws.cell(1, len(existing) + 1, header)
                existing.append(header)
    ws.freeze_panes = "A2"
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="20204B")
        cell.alignment = Alignment(horizontal="center")
    for i, width in enumerate([24, 26, 19, 22, 16, 28, 24, 20, 18, 18, 25], 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    return ws

def _serialize(wb):
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()

def _safe_text(value):
    value = str(value or "").strip()
    if value.startswith(("=", "+", "-", "@", "\t", "\r", "\n")):
        value = "'" + value
    return value

def _append(wb, student):
    event = student["event"]
    if event not in EVENTS:
        return None, "Invalid event."
    ws = _ensure_sheet(wb, event)
    mobile = student["mobile"]
    # Allow repeat submissions, including the same mobile number and event.
    reg_id = "ZEN-" + uuid.uuid4().hex[:10].upper()
    now = datetime.now(ZoneInfo("Asia/Kolkata"))
    ws.append([reg_id, _safe_text(student["name"]), mobile, _safe_text(student["department"]),
               _safe_text(student["year"]), _safe_text(student["college"]), event,
               now.strftime("%Y-%m-%d"), now.strftime("%I:%M:%S %p"),
               _safe_text(student.get("category", "")), _safe_text(student.get("batch", ""))])
    ws.cell(ws.max_row, 3).number_format = "@"
    return reg_id, None

class ExcelStore:
    def __init__(self, connection_string="", container="zen-event", blob_name="college_fest_registrations.xlsx"):
        self.cloud = bool(connection_string)
        if self.cloud:
            from azure.storage.blob import BlobServiceClient
            self.client = BlobServiceClient.from_connection_string(connection_string).get_blob_client(container=container, blob=blob_name)
        else:
            self.path = LOCAL_PATH

    def _load_cloud(self):
        from azure.core.exceptions import ResourceNotFoundError
        try:
            download = self.client.download_blob()
            return load_workbook(io.BytesIO(download.readall())), download.properties.etag
        except ResourceNotFoundError:
            return _new_workbook(), None

    def read_bytes(self):
        if self.cloud:
            wb, _ = self._load_cloud()
            for event in EVENTS:
                _ensure_sheet(wb, event)
            return _serialize(wb)
        with _LOCK:
            wb = load_workbook(self.path) if os.path.exists(self.path) else _new_workbook()
            for event in EVENTS:
                _ensure_sheet(wb, event)
            return _serialize(wb)

    def rows(self):
        wb = load_workbook(io.BytesIO(self.read_bytes()), read_only=True, data_only=True)
        result = []
        for ws in wb.worksheets:
            if ws.title not in EVENTS and ws.title != "Registrations":
                continue
            for row in ws.iter_rows(min_row=2, values_only=True):
                if row and row[0] is not None:
                    sheet_headers = [cell.value for cell in ws[1]]
                    result.append({header: dict(zip(sheet_headers, row)).get(header, "") for header in HEADERS})
        wb.close()
        return result

    def register(self, student):
        if not re.fullmatch(r"[6-9]\d{9}", student.get("mobile", "")):
            return None, "Enter a valid 10-digit Indian mobile number."
        if not 2 <= len(student.get("name", "").strip()) <= 100:
            return None, "Enter a valid student name."
        if student.get("event") not in EVENTS:
            return None, "Select a valid event."
        if not 2 <= len(student.get("batch", "").strip()) <= 80:
            return None, "Please enter your batch."
        if student["event"] in ("Singing", "Dance") and student.get("category") not in ("Single", "Duet", "Mix"):
            return None, "Select Single, Duet or Mix category."
        if self.cloud:
            from azure.core import MatchConditions
            from azure.core.exceptions import ResourceModifiedError, ResourceExistsError, HttpResponseError
            for attempt in range(8):
                wb, etag = self._load_cloud()
                reg_id, error = _append(wb, student)
                if error:
                    return None, error
                try:
                    if etag:
                        self.client.upload_blob(_serialize(wb), overwrite=True, etag=etag,
                                                match_condition=MatchConditions.IfNotModified)
                    else:
                        self.client.upload_blob(_serialize(wb), overwrite=False)
                    return reg_id, None
                except (ResourceModifiedError, ResourceExistsError):
                    time.sleep(0.1 * (attempt + 1))
                except HttpResponseError as exc:
                    if exc.status_code not in (409, 412):
                        raise
                    time.sleep(0.1 * (attempt + 1))
            return None, "High registration traffic. Please retry."
        with _LOCK:
            wb = load_workbook(self.path) if os.path.exists(self.path) else _new_workbook()
            reg_id, error = _append(wb, student)
            if error:
                return None, error
            temp = f"{self.path}.{uuid.uuid4().hex}.tmp"
            try:
                os.makedirs(os.path.dirname(os.path.abspath(self.path)), exist_ok=True)
                with open(temp, "wb") as file:
                    file.write(_serialize(wb))
                os.replace(temp, self.path)
            finally:
                if os.path.exists(temp):
                    os.remove(temp)
            return reg_id, None
